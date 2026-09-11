# -*- coding: utf-8 -*-
"""부록 계산근거 엑셀 — 교대·교각·상부 응력 계산의 입력·중간값·결과와 풀이 문장을 시트별로 남긴다. 출력 report/부록_계산근거_5장.xlsx"""
import sys, os, json
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
PJ = r"D:\Midas\projects\순천만IC2교"; RUNS = os.path.join(PJ, "runs"); sys.path.insert(0, PJ)
from derive import abutment_lines, pier_lines
A = json.load(open(os.path.join(RUNS, "abutment", "A1_v3_result.json"), encoding="utf-8"))
PR = {n: json.load(open(os.path.join(RUNS, "pier3", f"{n}_result.json"), encoding="utf-8")) for n in ("P1", "P2", "P3", "P4", "P5")}
GOV = json.load(open(os.path.join(RUNS, "stress3_gov.json"), encoding="utf-8")); RX = json.load(open(os.path.join(RUNS, "reactions_summary.json"), encoding="utf-8"))["support"]
CB = json.load(open(os.path.join(PJ, "combos_2010.json"), encoding="utf-8"))
HDR = PatternFill("solid", fgColor="DDDDDD"); B = Font(bold=True); TH = Side(style="thin"); BD = Border(left=TH, right=TH, top=TH, bottom=TH)

def sheet(wb, name, title, header, rows, widths=None, notes=None):
    ws = wb.create_sheet(name[:31]); ws["A1"] = title; ws["A1"].font = Font(bold=True, size=12); r0 = 3
    if notes:
        for k, ln in enumerate(notes): ws.cell(row=r0 + k, column=1, value=ln)
        r0 += len(notes) + 1
    for j, h in enumerate(header, 1): c = ws.cell(row=r0, column=j, value=h); c.font = B; c.fill = HDR; c.border = BD; c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r0 + i, column=j, value=v); c.border = BD
            if isinstance(v, float): c.number_format = "#,##0.000" if abs(v) < 10 else "#,##0.0"
    for j in range(1, len(header) + 1): ws.column_dimensions[get_column_letter(j)].width = (widths[j - 1] if widths else 14)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1); return ws

def main():
    wb = Workbook(); ws = wb.active; ws.title = "목차"
    ws["A1"] = "부록 — 5장 안전성평가 계산 근거 (자동 생성: abutment3.py, pier_model3.py, stress3.py 결과)"; ws["A1"].font = Font(bold=True, size=12)
    toc = ["교대_조건", "교대_하중집계", "교대_말뚝", "교대_단면", "교대_풀이", "교각_상부반력", "교각_하중별단면력", "교각_조합별단면력", "교각_코핑기둥기초", "교각_풀이", "상부_응력케이스", "상부_항복검토", "하중조합"]
    for i, t in enumerate(toc, 3): ws.cell(row=i, column=1, value=t)
    g, c, soil, pl, s = A["geom"], A["cond"], A["soil"], A["piles"], A["section"]; LA = abutment_lines(A)
    sheet(wb, "교대_조건", "교대 A1 검토조건", ["항목", "값", "단위", "출처"],
          [["앞굽", g["toe"], "m", "도면 자동판독"], ["벽체", g["stem"], "m", ""], ["뒷굽", g["heel"], "m", ""], ["B", g["B"], "m", ""], ["기초두께", g["tf"], "m", ""], ["받침면 높이", g["z_seat"], "m", ""], ["H", g["z_top"], "m", ""], ["흉벽 두께", g["t_par"], "m", ""], ["흉벽 높이", g["Hp"], "m", ""], ["폭", g["LW"], "m", ""],
           ["φ", soil["phi"], "°", "요령 표 2.7"], ["γ", soil["gamma"], "kN/m³", ""], ["q", soil["q"], "kN/m²", ""], ["Ka 안정(Rankine)", c["Ka_stab"], "", ""], ["Ka 벽체(Coulomb δ=φ/3)", c["Ka_wall"], "", ""], ["R_D 단위폭", c["D"], "kN/m", "격자해석"], ["R_L 단위폭", c["LL"], "kN/m", ""], ["받침 마찰계수", c["mu"], "", "편람 4-1"]], widths=[24, 12, 10, 24])
    rows = [[r["name"], r["kind"], r["V"], r["H"], r["x"], r["y"], r["Mr"], r["Mo"]] for r in A["loads"]] + [[k, "합계", v["V"], v["H"], v["xi"], v["e"], v["Mr"], v["Mo"]] for k, v in A["cases"].items()]
    sheet(wb, "교대_하중집계", "교대 A1 하중집계 (단위폭, 앞굽 연단 기준; 합계 행의 x = ξ, y = e)", ["구분", "종류", "V (kN)", "H (kN)", "x (m)", "y (m)", "Mr (kN·m)", "Mo (kN·m)"], rows, widths=[30, 8, 12, 12, 10, 10, 12, 12])
    keys = list(next(iter(pl["cases"].values()))["R"].keys())
    sheet(wb, "교대_말뚝", f"교대 A1 말뚝 반력 (강체 캡, 전폭; Ae = {pl['Ap']*1e6:.0f} mm², Ra = {pl['Ra']:.0f} kN)", ["조합", "ΣV", "ΣH", "M", "e"] + keys + ["수평/본", "f (MPa)", "v (MPa)"],
          [[k, v["V"], v["H"], v["M"], v["e"]] + list(v["R"].values()) + [v["Hp"], v["f"], v["v"]] for k, v in pl["cases"].items()], widths=[24] + [12] * (8 + len(keys)))
    st, pa, ft = s["stem"], s["parapet"], s["footing"]
    rows = [["벽체", r["name"], r["formula"], r["V"], r["arm"], r["M"], r["f"]] for r in st["rows"]] + [["흉벽", r["name"], r["formula"], r["V"], r["arm"], r["M"], r["f"]] for r in pa["rows"]]
    rows += [["벽체", "Mu/Vu", "", st["Vu"], "", st["Mu"], ""], ["흉벽 A(토압+상재)", "Mu/Vu", "", pa["VuA"], "", pa["MuA"], ""], ["흉벽 B(토압+윤하중)", "Mu/Vu", "", pa["Vu"], "", pa["Mu"], ""], ["앞굽", "Mu/Vu", "", ft["toe"]["Vu"], "", ft["toe"]["Mu"], ""], ["뒷굽", "Mu/Vu", "", ft["heel"]["Vu"], "", ft["heel"]["Mu"], ""],
             ["벽체", "φMn/φVc", st["bar"], st["phiVc"], st["d"], st["phiMn"], st["ratio_M"]], ["흉벽", "φMn/φVc", pa["bar"], pa["phiVc"], pa["d"], pa["phiMn"], pa["ratio_M"]], ["기초", "φMn/φVc", ft["bar"], ft["phiVc"], ft["d"], ft["phiMn"], max(ft["toe"]["ratio_M"], ft["heel"]["ratio_M"])]]
    sheet(wb, "교대_단면", "교대 A1 단면검토 (강도 ①)", ["부재", "하중/항목", "산정식·배근", "V 또는 φVc (kN/m)", "팔길이·d", "M 또는 φMn (kN·m/m)", "계수·비"], rows, widths=[18, 20, 30, 16, 12, 18, 10])
    sheet(wb, "교대_풀이", "교대 A1 풀이과정", ["단계", "풀이"], [[k, ln] for k, v in LA.items() for ln in v], widths=[10, 160])
    rows = []
    for n, rx in RX.items():
        if n == "A1": continue
        L1 = rx["L_one"]; rows.append([n, rx["D_g"]["G1"], rx["D_g"]["G2"], rx["D"], rx["L_g"]["G1"][0], rx["L_g"]["G2"][0], rx["Lmax"], L1["L1"]["G1"], L1["L1"]["G2"], L1["L2"]["G1"], L1["L2"]["G2"], rx["H"]["W"][1], rx["H"]["WL"][1], rx["H"]["CF"][1], rx["H"]["LF"][0], rx["H"]["TP"][0]])
    sheet(wb, "교각_상부반력", "교각 상부반력 (격자해석, kN; 받침 국부축)", ["교각", "D G1", "D G2", "D 계", "Lmax G1", "Lmax G2", "Lmax 계", "편재 L1 G1", "편재 L1 G2", "편재 L2 G1", "편재 L2 G2", "W(y)", "WL(y)", "CF(y)", "LF(x)", "TP(x)"], rows, widths=[8] + [11] * 15)
    rows = []
    for n, r in PR.items():
        for k, v in r["per_load"].items(): b = v["base"]; t = v["top"]; rows.append([n, k, b["P"], b["Vx"], b["Vy"], b["Mx"], b["My"], t["P"], t["Mx"], t["My"]])
    sheet(wb, "교각_하중별단면력", "교각 하중별 기둥 하단(기초 상면)·상단 단면력 (NX)", ["교각", "하중", "P", "Vx", "Vy", "Mx", "My", "상단 P", "상단 Mx", "상단 My"], rows, widths=[8, 10] + [12] * 8)
    rows = []
    for n, r in PR.items():
        for cmb in r["combos"]: rows.append([n, cmb["key"], cmb["kind"], cmb["fD"], cmb["P"], cmb["Vx"], cmb["Vy"], cmb["Mx"], cmb["My"], cmb["Pf"], cmb["Mxf"], cmb["Myf"], cmb["pile_max"], cmb["pile_min"], cmb["pile_h"]])
    sheet(wb, "교각_조합별단면력", "교각 조합별 기둥 하단 단면력·기초 저면·말뚝 반력 (전 조합)", ["교각", "조합", "종류", "D계수", "P", "Vx", "Vy", "Mx", "My", "Pf(저면)", "Mxf", "Myf", "말뚝 max", "말뚝 min", "수평/본"], rows, widths=[8, 30, 8, 8] + [12] * 11)
    rows = []
    for n, r in PR.items():
        c, w, f, b = r["coping"], r["column"]["worst"], r["footing"], r["bearing"]
        rows.append([n, c["av"], c["d"], c["Vu"], c["Mu"], c["As_req"], c["As_min"], c["As_use"], c["ratio_As"], w["key"], w["P"], w["M"], w["ds"], w["ratio"], r["column"]["shear"]["ratio"], f["Kv"], f["beta_lam"], f["pile_max"], f["Ra"], f["ratio_pile"], f["section"]["Mu"], f["section"]["phiMn"], f["section"]["ratio_M"], b["R"], b["cap"], b["ratio"]])
    sheet(wb, "교각_코핑기둥기초", "교각 부재검토 요약", ["교각", "av", "d", "코핑 Vu", "코핑 Mu", "As,req", "As,min", "As 사용", "As 비", "기둥 조합", "Pu", "Mu", "δs", "P–M 비", "전단비", "Kv", "βλ", "말뚝 max", "Ra", "말뚝비", "기초 Mu", "φMn", "기초비", "받침 반력", "받침 용량", "받침비"], rows, widths=[8] + [11] * 25)
    sheet(wb, "교각_풀이", "교각 풀이과정 (P1~P5)", ["교각", "항목", "풀이"], [[n, k, ln] for n, r in PR.items() for k, v in pier_lines(r).items() for ln in v], widths=[8, 10, 160])
    rows = []
    for tag in ("pos", "neg"):
        for cid, cc in GOV[tag]["cases"].items():
            rows.append([tag, cid, cc["max"][0], cc["min"][0], cc["max"][1], cc["min"][1], cc["max"][2], cc["max"][4], cc["fa"], cc["fca_top"], cc["fca_bot"], cc["fc"], cc.get("fac_c"), cc.get("fac_t"), cc["sf"]])
    sheet(wb, "상부_응력케이스", "상부 거더 응력 케이스 (정모멘트부 pos / 부모멘트부 neg, MPa)", ["단면", "case", "강재 상연 max", "min", "강재 하연 max", "min", "바닥판 상연", "바닥판 하연", "fa", "fca 상", "fca 하", "fc", "증가 압축", "증가 인장", "S.F"], rows, widths=[8, 6] + [11] * 13)
    rows = [[tag, y["fiber"], y["sign"], y["D13"], y["L215"], y["creep"], y["shrink"], y["temp"], y["sum"], y["fy"], y["ratio"]] for tag in ("pos", "neg") for y in GOV["yield_check"][tag]]
    sheet(wb, "상부_항복검토", "항복에 대한 안전도 검사 (3.9.3.2)", ["단면", "연", "부호", "1.3D", "2.15L", "크리프", "건조수축", "온도차", "Σf", "허용", "비"], rows, widths=[8, 14, 6] + [10] * 8)
    rows = [["강도 " + k, f["식"], f["D"], f["L"], f["CF"], f["H"], f["W"], f["WL"], f["BK"], f["G"], f.get("D_min")] for k, f in CB["강도"].items()] + [["사용 " + k, f["식"], f["D"], f["L"], f["CF"], f["H"], f["W"], f["WL"], f["BK"], f["G"], f.get("허용증가")] for k, f in CB["사용"].items()]
    sheet(wb, "하중조합", "하중조합 (도로교설계기준 2010; 지진 ⑦·사용 6은 보고서에서 제외)", ["조합", "식", "D", "L", "CF", "H", "W", "WL", "BK", "G", "D_min/증가"], rows, widths=[10, 50] + [7] * 9)
    out = os.path.join(PJ, "report", "부록_계산근거_5장.xlsx"); wb.save(out); print("저장", out)

if __name__ == "__main__":
    main()
